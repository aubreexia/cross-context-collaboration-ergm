#!/usr/bin/env python3
"""Export selected aggregate sheets from an analysis workbook as public CSVs.

The supplied ERGM workbooks contain absolute local paths in some summary
columns. This utility reads a workbook in read-only mode, drops path-like
columns, and writes only requested aggregate sheets. It does not anonymize
actor-level data; do not use it on node, edge, or dyad-level sheets.
"""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

from openpyxl import load_workbook


PATH_LIKE_COLUMNS = re.compile(r"(?:^|_)(?:file|path|directory|root)(?:$|_)", re.I)
DISALLOWED_SHEETS = re.compile(
    r"(?:nodes|edges|dyads|matrix|transformed_nodes|combined_nodes|combined_edges)",
    re.I,
)


def safe_name(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("_") or "sheet"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument(
        "--sheets",
        required=True,
        help="Comma-separated worksheet names to export.",
    )
    return parser.parse_args()


def contains_unsafe_path(value: object) -> bool:
    if not isinstance(value, str):
        return False
    return bool(re.search(r"(?:^|\s)(?:[A-Za-z]:[\\/]|/Users/|/groups/|/home/)", value))


def main() -> None:
    args = parse_args()
    source = args.input.expanduser().resolve()
    destination = args.output_dir.expanduser().resolve()
    requested = [name.strip() for name in args.sheets.split(",") if name.strip()]
    if not source.is_file():
        raise FileNotFoundError(source)
    if not requested:
        raise ValueError("Specify at least one worksheet with --sheets.")

    workbook = load_workbook(source, read_only=True, data_only=True)
    unavailable = sorted(set(requested) - set(workbook.sheetnames))
    if unavailable:
        raise ValueError(
            f"Worksheet(s) not found: {', '.join(unavailable)}. "
            f"Available: {', '.join(workbook.sheetnames)}"
        )

    destination.mkdir(parents=True, exist_ok=True)
    for sheet_name in requested:
        if DISALLOWED_SHEETS.search(sheet_name):
            raise ValueError(
                f"Refusing to export potentially actor/dyad-level sheet: {sheet_name}"
            )
        worksheet = workbook[sheet_name]
        # Some supplied workbooks have an incomplete XML dimension; this makes
        # openpyxl iterate all populated rows rather than silently truncating.
        worksheet.reset_dimensions()
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            continue
        header_text = ["" if value is None else str(value).strip() for value in header]
        retain = [
            index
            for index, name in enumerate(header_text)
            if name and not PATH_LIKE_COLUMNS.search(name)
        ]
        output = destination / f"{safe_name(sheet_name)}.csv"
        with output.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow([header_text[index] for index in retain])
            for row in rows:
                values = []
                for index in retain:
                    value = row[index] if index < len(row) else None
                    values.append("[path removed]" if contains_unsafe_path(value) else value)
                writer.writerow(values)
        print(f"[SAVED] {output}")


if __name__ == "__main__":
    main()
